import csv
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from app.importers.phase2 import extract, reconcile, run


class Phase2Tests(unittest.TestCase):
    def test_csv_reconciles_every_url_and_preserves_source(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.csv"
            rows = [
                ["1", "https://www.instagram.com/reel/ABC/?igsh=one", "first"],
                ["2", "https://www.instagram.com/reel/ABC/?igsh=one", "exact copy"],
                ["3", "https://www.instagram.com/reel/ABC/?igsh=two", "variant"],
                ["4", "https:///broken", "bad"],
                ["5", "https://example.com/article", "website"],
            ]
            with source.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(["Record #", "URL", "Description"])
                writer.writerows(rows)
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            report = run(source, root / "report")
            self.assertEqual(report["detected_url_occurrences"], 5)
            self.assertEqual(report["accounted_url_occurrences"], 5)
            self.assertTrue(report["reconciled"])
            self.assertEqual(report["buckets"], {"importable": 2, "exact_duplicates": 1, "invalid": 1, "needs_review": 1})
            self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())
            self.assertTrue((root / "report" / "reconciliation_report.json").exists())
            self.assertEqual(json.loads((root / "report" / "reconciliation_report.json").read_text())["source"]["sha256_before"], before)

    def test_docx_extracts_url_in_source_order(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.docx"
            document = Document()
            document.add_paragraph("A useful link https://example.com/one")
            document.add_paragraph("Compare https://example.com/two and https://example.com/three")
            document.save(source)
            records = extract(source)
            report = reconcile(records)
            self.assertEqual([record["original_url"] for record in records], ["https://example.com/one", "https://example.com/two", "https://example.com/three"])
            self.assertEqual(records[0]["description_confidence"], "CERTAIN")
            self.assertEqual(records[1]["description_confidence"], "NEEDS_REVIEW")
            self.assertTrue(report["reconciled"])

    def test_refuses_to_overwrite_existing_report(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.csv"
            source.write_text("URL,Description\nhttps://example.com,ok\n", encoding="utf-8")
            output = root / "report"
            run(source, output)
            with self.assertRaises(FileExistsError):
                run(source, output)


if __name__ == "__main__":
    unittest.main()
