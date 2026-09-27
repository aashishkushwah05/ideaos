"""Phase 2: deterministic source extraction and zero-loss reconciliation.

This module reads CSV/DOCX inputs only. It does not use AI or SQLite and never
modifies a source file. Normalized values are comparison metadata only.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from docx import Document

URL_PATTERN = re.compile(r'https?://[^\s<>"\'\]\[(){}]+', re.IGNORECASE)
TRACKING_KEYS = {"igsh", "igshid", "fbclid", "gclid", "utm_source", "utm_medium", "utm_campaign", "utm_term", "utm_content"}
STATUSES = {"IMPORTABLE", "EXACT_DUPLICATE", "INVALID", "NEEDS_REVIEW"}


def find_urls(text: str) -> list[str]:
    return [m.group(0).rstrip(".,;:!?") for m in URL_PATTERN.finditer(text or "")]


def inspect_url(original_url: str) -> dict[str, str | None] | None:
    """Validate HTTP(S) URL and return comparison-only metadata."""
    try:
        parts = urlsplit(original_url)
        host = (parts.hostname or "").lower()
        port = parts.port
    except ValueError:
        return None
    if parts.scheme.lower() not in {"http", "https"} or not host:
        return None
    if host.startswith("www."):
        host = host[4:]
    netloc = host
    if port and not ((parts.scheme.lower(), port) in {("http", 80), ("https", 443)}):
        netloc += f":{port}"
    path = re.sub(r"/{2,}", "/", parts.path or "/")
    if path != "/":
        path = path.rstrip("/")
    query = urlencode(sorted((k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True) if k.lower() not in TRACKING_KEYS))
    platform, content_id = "Website", None
    bits = [part for part in path.split("/") if part]
    if host in {"instagram.com", "m.instagram.com"}:
        platform = "Instagram"
        if len(bits) >= 2 and bits[0].lower() in {"reel", "reels", "p", "tv"}:
            content_id = f"{bits[0].lower()}:{bits[1]}"
    elif host in {"youtube.com", "m.youtube.com", "youtu.be"}:
        platform = "YouTube"
        if host == "youtu.be" and bits:
            content_id = bits[0]
    elif host in {"github.com", "gist.github.com"}:
        platform = "GitHub"
    elif host in {"x.com", "twitter.com", "mobile.twitter.com"}:
        platform = "X/Twitter"
    elif host in {"linkedin.com", "in.linkedin.com"}:
        platform = "LinkedIn"
    elif host in {"facebook.com", "m.facebook.com", "fb.watch"}:
        platform = "Facebook"
    return {"normalized_url": urlunsplit((parts.scheme.lower(), netloc, path, query, "")), "platform": platform, "platform_content_id": content_id}


def _record(records: list[dict], url: str, description: str | None, location: str, confidence: str, source_number: str | None = None) -> None:
    records.append({"record_number": len(records) + 1, "source_record_number": source_number, "original_url": url, "original_description": description, "source_location": location, "description_confidence": confidence, "normalized_url": None, "platform": None, "platform_content_id": None, "status": None, "reasons": [], "exact_duplicate_of": None, "possible_duplicate_of": []})


def extract_csv(path: Path) -> list[dict]:
    records: list[dict] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError("CSV has no header row.")
        names = {name.strip().lower(): name for name in reader.fieldnames if name}
        url_key = next((names[k] for k in ("url", "original_url", "link") if k in names), None)
        description_key = next((names[k] for k in ("description", "original_description", "notes", "note") if k in names), None)
        number_key = next((names[k] for k in ("record #", "record_number", "record") if k in names), None)
        for row_number, row in enumerate(reader, 2):
            occurrences = [(url, key == url_key) for key, value in row.items() for url in find_urls(value or "")]
            for index, (url, from_url_column) in enumerate(occurrences, 1):
                description = (row.get(description_key) or "").strip() if description_key else ""
                certain = bool(url_key and description_key and from_url_column and len(occurrences) == 1)
                source_number = (row.get(number_key) or "").strip() if number_key else None
                _record(records, url, description or None, f"csv:row={row_number},url_occurrence={index}", "CERTAIN" if certain else "NEEDS_REVIEW", source_number or None)
    return records


def extract_docx(path: Path) -> list[dict]:
    document, records = Document(path), []
    elements = [(f"docx:paragraph={i}", p.text) for i, p in enumerate(document.paragraphs, 1)]
    for ti, table in enumerate(document.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                elements.append((f"docx:table={ti},row={ri},cell={ci}", cell.text))
    for location, text in elements:
        urls = find_urls(text)
        description = " ".join(re.sub(URL_PATTERN, "", text).split()) or None
        for url in urls:
            _record(records, url, description, location, "CERTAIN" if len(urls) == 1 else "NEEDS_REVIEW")
    return records


def extract(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        return extract_csv(path)
    if path.suffix.lower() == ".docx":
        return extract_docx(path)
    raise ValueError("Unsupported source type. Use a .csv or .docx file.")


def reconcile(records: list[dict]) -> dict:
    exact, variants = {}, {}
    for record in records:
        details = inspect_url(record["original_url"])
        if details is None:
            record["status"] = "INVALID"
            record["reasons"].append("URL failed deterministic HTTP(S) validation.")
            continue
        record.update(details)
        original = record["original_url"]
        if original in exact:
            record["status"] = "EXACT_DUPLICATE"
            record["exact_duplicate_of"] = exact[original]
            record["reasons"].append(f"Exact original URL duplicate of record {exact[original]}.")
            continue
        exact[original] = record["record_number"]
        key = (record["platform"], record["platform_content_id"]) if record["platform_content_id"] else None
        if key and key in variants:
            record["status"] = "NEEDS_REVIEW"
            record["possible_duplicate_of"] = list(variants[key])
            record["reasons"].append("Same platform content ID as another original URL; not auto-merged.")
        elif record["description_confidence"] != "CERTAIN":
            record["status"] = "NEEDS_REVIEW"
            record["reasons"].append("URL-to-description relationship is not deterministically certain.")
        else:
            record["status"] = "IMPORTABLE"
        if key:
            variants.setdefault(key, []).append(record["record_number"])
    counts = Counter(record["status"] for record in records)
    buckets = {"importable": counts["IMPORTABLE"], "exact_duplicates": counts["EXACT_DUPLICATE"], "invalid": counts["INVALID"], "needs_review": counts["NEEDS_REVIEW"]}
    return {"detected_url_occurrences": len(records), "accounted_url_occurrences": sum(buckets.values()), "reconciled": len(records) == sum(buckets.values()), "buckets": buckets, "unique_original_url_values": len(exact), "records": records}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def run(source: Path, output: Path) -> dict:
    source, output = source.resolve(strict=True), output.resolve()
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"Refusing to overwrite non-empty output directory: {output}")
    if output in source.parents or output == source.parent:
        raise ValueError("Output directory cannot contain the source file.")
    before = sha256(source)
    report = reconcile(extract(source))
    after = sha256(source)
    report["source"] = {"path": str(source), "sha256_before": before, "sha256_after": after, "unchanged_during_extraction": before == after}
    report["generated_at_utc"] = datetime.now(UTC).isoformat()
    report["phase"] = "Phase 2 — deterministic extraction and zero-loss reconciliation"
    output.mkdir(parents=True, exist_ok=False)
    (output / "reconciliation_report.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    with (output / "extracted_records.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=report["records"][0].keys() if report["records"] else ["record_number"])
        writer.writeheader(); writer.writerows(report["records"])
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="AI Idea Vault Phase 2 deterministic extractor")
    parser.add_argument("source", type=Path); parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(); report = run(args.source, args.output)
    print(json.dumps({key: report[key] for key in ("detected_url_occurrences", "accounted_url_occurrences", "reconciled", "buckets")}, indent=2))
    return 0 if report["reconciled"] and report["source"]["unchanged_during_extraction"] else 1

if __name__ == "__main__":
    raise SystemExit(main())
